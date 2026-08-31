from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path('artifacts/client-documents/CR-Mariposa-Website-CMS-Guide.docx')

# compact_reference_guide preset plus one named brand override:
# CR Mariposa's sand, basalt, clay, and moss palette replaces the preset blues.
SAND = 'EAE3D6'
SAND_LIGHT = 'F6F2EA'
BASALT = '1D1F1C'
CLAY = '9C5F3C'
CLAY_DARK = '7A4526'
MOSS = '4A4F46'
STONE = '6E6A5F'
WHITE = 'FFFFFF'
HAIRLINE = 'D3CDC1'


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name='Calibri', size=11, color=BASALT, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=120, start=180, bottom=120, end=180):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(total))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent_dxa))
    tbl_ind.set(qn('w:type'), 'dxa')

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement('w:gridCol')
        grid_col.set(qn('w:w'), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(width))
            tc_w.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color=HAIRLINE, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    tr_pr.append(repeat)


def add_hyperlink(paragraph, text, url, color=CLAY_DARK):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship_id)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    run_fonts = OxmlElement('w:rFonts')
    run_fonts.set(qn('w:ascii'), 'Calibri')
    run_fonts.set(qn('w:hAnsi'), 'Calibri')
    props.append(run_fonts)
    tint = OxmlElement('w:color')
    tint.set(qn('w:val'), color)
    props.append(tint)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    text_node = OxmlElement('w:t')
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t')
    value.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, value, end])
    set_run_font(run, size=9, color=STONE)


def add_numbering_definition(doc, numbered=False):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
    abstract_id = max(existing_abstract, default=-1) + 1
    existing_nums = [int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]
    num_id = max(existing_nums, default=0) + 1

    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    multi = OxmlElement('w:multiLevelType')
    multi.set(qn('w:val'), 'singleLevel')
    abstract.append(multi)
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:start')
    start.set(qn('w:val'), '1')
    level.append(start)
    num_fmt = OxmlElement('w:numFmt')
    num_fmt.set(qn('w:val'), 'decimal' if numbered else 'bullet')
    level.append(num_fmt)
    level_text = OxmlElement('w:lvlText')
    level_text.set(qn('w:val'), '%1.' if numbered else '•')
    level.append(level_text)
    justification = OxmlElement('w:lvlJc')
    justification.set(qn('w:val'), 'left')
    level.append(justification)
    p_pr = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'num')
    tab.set(qn('w:pos'), '540')
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement('w:ind')
    indent.set(qn('w:left'), '540')
    indent.set(qn('w:hanging'), '270')
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    first_num = numbering.find(qn('w:num'))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    abstract_ref = OxmlElement('w:abstractNumId')
    abstract_ref.set(qn('w:val'), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_list_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn('w:numPr'))
    if num_pr is None:
        num_pr = OxmlElement('w:numPr')
        p_pr.append(num_pr)
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    num_id_el = OxmlElement('w:numId')
    num_id_el.set(qn('w:val'), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_bullet(doc, text, bullet_id, bold_lead=None):
    p = doc.add_paragraph()
    apply_list_numbering(p, bullet_id)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        tail = p.add_run(text[len(bold_lead):])
        set_run_font(tail)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_step(doc, number_id, heading, body):
    p = doc.add_paragraph()
    apply_list_numbering(p, number_id)
    lead = p.add_run(f'{heading}. ')
    set_run_font(lead, bold=True, color=CLAY_DARK)
    rest = p.add_run(body)
    set_run_font(rest)
    return p


def add_callout(doc, label, body, fill=SAND_LIGHT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=HAIRLINE)
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    first = p.add_run(f'{label}  ')
    set_run_font(first, size=10.5, color=CLAY_DARK, bold=True)
    second = p.add_run(body)
    set_run_font(second, size=10.5, color=BASALT)


def add_section_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11.5, color=MOSS)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BASALT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ('Heading 1', 16, CLAY_DARK, 18, 10),
        ('Heading 2', 13, CLAY, 14, 7),
        ('Heading 3', 12, BASALT, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = 'Georgia'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Georgia')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Georgia')
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run('CR MARIPOSA RENTALS')
    set_run_font(left, size=8.5, color=STONE, bold=True)
    right = p.add_run('\tWEBSITE CONTENT GUIDE')
    set_run_font(right, size=8.5, color=STONE)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    note = fp.add_run('Owner and client editor reference')
    set_run_font(note, size=9, color=STONE)
    page_text = fp.add_run('\tPage ')
    set_run_font(page_text, size=9, color=STONE)
    add_page_field(fp)


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run('CLIENT ENABLEMENT GUIDE')
    set_run_font(run, size=10, color=CLAY, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 0.95
    run = title.add_run('Website Content\nManagement Guide')
    set_run_font(run, name='Georgia', size=31, color=BASALT)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run(
        'How CR Mariposa owners and client editors keep listings, photography, '
        'guest information, and bilingual content current - without touching code.'
    )
    set_run_font(run, size=13.5, color=MOSS)

    table = doc.add_table(rows=2, cols=2)
    set_table_geometry(table, [4680, 4680])
    set_table_borders(table, color=HAIRLINE)
    entries = [
        ('WEBSITE', 'cr-mariposa-staging.netlify.app'),
        ('ADMIN', 'cr-mariposa-staging.netlify.app/admin'),
        ('CONTENT', 'English and Spanish'),
        ('VERSION', 'Owner guide 1.0 | August 2026'),
    ]
    for index, (label, value) in enumerate(entries):
        row, column = divmod(index, 2)
        cell = table.cell(row, column)
        set_cell_fill(cell, SAND_LIGHT if row == 0 else WHITE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(label)
        set_run_font(label_run, size=8.5, color=CLAY, bold=True)
        value_p = cell.add_paragraph()
        value_p.paragraph_format.space_after = Pt(0)
        value_run = value_p.add_run(value)
        set_run_font(value_run, size=10.5, color=BASALT)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    add_callout(
        doc,
        'CLIENT CONTROL',
        'Content and publishing are editable. Layout, visual design, filters, security, and application behavior remain protected.',
        SAND,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run('Prepared for CR Mariposa Rentals')
    set_run_font(run, size=10.5, color=STONE, italic=True)


def add_start_page(doc, bullet_id, number_id):
    heading = doc.add_heading('Start here', level=1)
    heading.paragraph_format.page_break_before = True
    add_section_intro(
        doc,
        'The content manager is designed for everyday ownership. A client editor can keep the public website accurate while the underlying design and technical system stay protected.',
    )

    doc.add_heading('Sign in', level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run('Admin address: '), bold=True)
    add_hyperlink(p, 'Open the CR Mariposa admin', 'https://cr-mariposa-staging.netlify.app/admin')
    add_bullet(doc, 'Use your assigned email and password. Credentials should never be shared in a presentation or placed in this guide.', bullet_id)
    add_bullet(doc, 'Use an Editor account for normal client access. Administrator access is reserved for account management and technical ownership.', bullet_id)

    doc.add_heading('The safe publishing workflow', level=2)
    add_step(doc, number_id, 'Choose the language', 'Select English or Spanish before editing a localized field.')
    add_step(doc, number_id, 'Open the content area', 'Choose Properties, Media, Reviews, Inquiries, a Page, or shared Settings.')
    add_step(doc, number_id, 'Make the change', 'Update only information that has been confirmed by the owner.')
    add_step(doc, number_id, 'Save as draft', 'Draft changes remain private; the last published version stays live.')
    add_step(doc, number_id, 'Publish when approved', 'Refresh the relevant public page and confirm the result on desktop and mobile.')

    add_callout(
        doc,
        'LANGUAGE REMINDER',
        'English and Spanish values are edited separately. Before publishing, check both languages whenever a public field is localized.',
    )

    doc.add_heading('What changes immediately?', level=2)
    add_bullet(doc, 'Published property facts, photographs, reviews, page copy, and contact details flow to the public website.', bullet_id)
    add_bullet(doc, 'Drafts do not replace the current public version.', bullet_id)
    add_bullet(doc, 'Unpublishing a property removes it from public property lists and detail pages.', bullet_id)


def add_properties_page(doc, bullet_id):
    heading = doc.add_heading('Properties and photography', level=1)
    heading.paragraph_format.page_break_before = True
    add_section_intro(
        doc,
        'Each property record controls its card, detail page, photo experience, key facts, and selected homepage placement. This is the area clients will use most often.',
    )

    doc.add_heading('Property basics', level=2)
    for text, lead in [
        ('Name and URL: Edit the public name. Avoid changing the URL slug after launch.', 'Name and URL:'),
        ('Location: Maintain the region, public district, community name, and optional badge. Never enter a private street address.', 'Location:'),
        ('Description: Maintain short and full descriptions in English and Spanish.', 'Description:'),
        ('Display: Feature a property, set its order, and publish or unpublish it.', 'Display:'),
    ]:
        add_bullet(doc, text, bullet_id, lead)

    doc.add_heading('Photos', level=2)
    for text, lead in [
        ('Hero image: The main photograph used on cards and the property page.', 'Hero image:'),
        ('Gallery order: Drag images into the intended viewing sequence.', 'Gallery order:'),
        ('Categories: Organize images as exterior, living room, kitchen, bedroom, amenities, or other.', 'Categories:'),
        ('Showcase selection: Choose the strongest cross-section for the opening photo mosaic.', 'Showcase selection:'),
        ('Alt text: Describe every image accurately in English and Spanish for accessibility and search.', 'Alt text:'),
    ]:
        add_bullet(doc, text, bullet_id, lead)

    add_callout(
        doc,
        'PHOTO STANDARD',
        'Lead with a photograph that unmistakably represents the specific home. Community or resort-style amenity images may remain in the gallery only when ownership and context are confirmed.',
    )

    doc.add_heading('Facts, amenities, and sleeping arrangements', level=2)
    add_bullet(doc, 'Bedrooms, beds, bathrooms, parking, and owner-confirmed maximum occupancy.', bullet_id)
    add_bullet(doc, 'Grouped amenities such as Wi-Fi, air conditioning, kitchen, laundry, pool, and security.', bullet_id)
    add_bullet(doc, 'Room-by-room “Where you’ll sleep” cards with a representative photo and verified bed description.', bullet_id)
    add_bullet(doc, 'Cancellation wording, check-in/out times, smoking, pets, events, and minimum-stay information.', bullet_id)
    add_bullet(doc, 'Approved external listing links, approximate map coordinates, and page-level SEO.', bullet_id)

    add_callout(
        doc,
        'DO NOT GUESS',
        'Leave ratings, bed sizes, occupancy, rules, or exact coordinates blank until the owner confirms them.',
        'FFF7ED',
    )


def add_site_content_page(doc, bullet_id):
    heading = doc.add_heading('Pages, settings, reviews, and inquiries', level=1)
    heading.paragraph_format.page_break_before = True
    add_section_intro(
        doc,
        'The remaining admin areas keep the brand message, contact channels, social proof, and inquiry workflow current across the entire site.',
    )

    doc.add_heading('Pages', level=2)
    add_bullet(doc, 'Home page: hero wording and images, property-row headings, review introduction, trust content, contact copy, and SEO.', bullet_id, 'Home page:')
    add_bullet(doc, 'Properties page: the All Homes heading, introduction, and SEO.', bullet_id, 'Properties page:')

    doc.add_heading('Shared settings', level=2)
    add_bullet(doc, 'Site settings: business name, tagline, email, telephone, WhatsApp number and default message, business location, social links, and default SEO.', bullet_id, 'Site settings:')
    add_bullet(doc, 'Rental settings: Airbnb, Booking.com, Vrbo, Expedia, and social-profile links plus direct-inquiry wording.', bullet_id, 'Rental settings:')
    add_bullet(doc, 'Map details: clients can maintain district-level information and approximate coordinates. Activating or changing the Google Maps API key remains a developer task.', bullet_id, 'Map details:')

    doc.add_heading('Reviews', level=2)
    add_bullet(doc, 'Publish only genuine reviews CR Mariposa is permitted to reuse. Add the original quote, approved translation, guest details, rating, platform, and optional property relationship.', bullet_id)
    add_bullet(doc, 'Feature selected reviews on the homepage and retain the source URL privately for provenance.', bullet_id)

    doc.add_heading('Guest inquiries', level=2)
    add_bullet(doc, 'View the guest’s name, contact details, message, requested dates, language, property, source, and submission time.', bullet_id)
    add_bullet(doc, 'Move each inquiry through New, Contacted, and Closed. WhatsApp and telephone clicks are direct actions and do not create an inquiry record.', bullet_id)

    doc.add_heading('Media library', level=2)
    add_bullet(doc, 'Upload new photographs, maintain captions and bilingual alt text, and reuse assets across properties.', bullet_id)
    add_bullet(doc, 'Removing an image from one gallery does not delete the media file. Delete media only after confirming it is not used anywhere else.', bullet_id)


def add_roles_page(doc, bullet_id):
    heading = doc.add_heading('Access, guardrails, and final checks', level=1)
    heading.paragraph_format.page_break_before = True
    add_section_intro(
        doc,
        'Client access is intentionally powerful for content and intentionally limited for design and technical behavior. This keeps routine ownership simple without risking the website system.',
    )

    doc.add_heading('Recommended access levels', level=2)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1700, 3730, 3930])
    set_table_borders(table)
    headers = ['ROLE', 'CAN MANAGE', 'RECOMMENDED FOR']
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_fill(cell, SAND)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(label), size=9, color=CLAY_DARK, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ('Editor', 'Properties, media, pages, reviews, settings, and inquiries', 'Owner, property manager, or client content team'),
        ('Administrator', 'Everything an Editor can do, plus user-account administration', 'Agency or designated system owner'),
        ('Developer', 'Layout, components, navigation behavior, integrations, security, and deployment', 'Technical implementation and support'),
    ]
    for role, scope, audience in rows:
        cells = table.add_row().cells
        for index, value in enumerate((role, scope, audience)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=9.5, color=BASALT, bold=(index == 0))
        set_table_geometry(table, [1700, 3730, 3930])

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.add_heading('Protected from direct client editing', level=2)
    add_bullet(doc, 'Page layout, fonts, colors, reusable components, filters, and responsive behavior.', bullet_id)
    add_bullet(doc, 'Navigation structure, application code, security rules, hosting, database, and Cloudflare R2 configuration.', bullet_id)
    add_bullet(doc, 'Booking calendars, payment processing, reservations, and platform integrations, which are outside the current website scope.', bullet_id)

    doc.add_heading('Before publishing', level=2)
    for text in [
        'Confirm the fact or policy with the owner, then check both English and Spanish.',
        'Use a property-specific hero photograph and meaningful alt text.',
        'Keep exact private addresses and sensitive information out of public fields.',
        'Check links, telephone numbers, and WhatsApp messages.',
        'Review the result on desktop and mobile after publishing.',
    ]:
        add_bullet(doc, text, bullet_id)

    add_callout(
        doc,
        'WHEN TO ASK FOR SUPPORT',
        'Contact the developer for layout changes, new page types, integrations, login or account problems, map-key configuration, deployment issues, or anything that changes how the website works rather than what it says.',
        SAND,
    )

def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    bullet_id = add_numbering_definition(doc, numbered=False)
    number_id = add_numbering_definition(doc, numbered=True)

    props = doc.core_properties
    props.title = 'CR Mariposa Website Content Management Guide'
    props.subject = 'Client-facing guide to the CR Mariposa Rentals content management system'
    props.author = 'CR Mariposa Rentals'
    props.keywords = 'CR Mariposa, CMS, Payload, property management, website guide'
    props.comments = 'Owner and client editor reference guide'

    add_cover(doc)
    add_start_page(doc, bullet_id, number_id)
    add_properties_page(doc, bullet_id)
    add_site_content_page(doc, bullet_id)
    add_roles_page(doc, bullet_id)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == '__main__':
    build_document()
